import os
import io
import uuid
from pathlib import Path
from typing import List, Dict, Any, Optional
from PIL import Image, ImageOps
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION, WD_ORIENTATION
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from ..config import DOCS_DIR, sanitize_filename

def fix_image_orientation_and_compress(image_path: Path, max_dimension: int = 2400) -> io.BytesIO:
    """
    Auto-corrects EXIF orientation, applies smart resizing if excessively large,
    and returns a clean BytesIO buffer for docx insertion.
    """
    with Image.open(image_path) as img:
        # Auto-rotate based on EXIF tag
        img = ImageOps.exif_transpose(img)
        
        # Convert RGBA or CMYK to RGB if saving as JPEG
        format_to_use = img.format if img.format in ("JPEG", "PNG") else "JPEG"
        if img.mode in ("RGBA", "P") and format_to_use == "JPEG":
            img = img.convert("RGB")
        elif img.mode not in ("RGB", "RGBA"):
            img = img.convert("RGB")
            format_to_use = "JPEG"

        # Resize if dimensions exceed max_dimension
        w, h = img.size
        if max(w, h) > max_dimension:
            scale = max_dimension / max(w, h)
            new_w = int(w * scale)
            new_h = int(h * scale)
            img = img.resize((new_w, new_h), Image.Resampling.LANCZOS)

        output = io.BytesIO()
        if format_to_use == "PNG":
            img.save(output, format="PNG", optimize=True)
        else:
            img.save(output, format="JPEG", quality=88, optimize=True)
        output.seek(0)
        return output

def add_footer_and_header(doc: Document, title: str):
    """Adds a modern header and subtle footer to the Word document."""
    for section in doc.sections:
        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run(f"Snap2Doc | {title[:40]}")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(148, 163, 184) # slate-400

        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Created with Snap2Doc — Modern Photo to Word Document Generator")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8)
        frun.font.color.rgb = RGBColor(148, 163, 184)

def generate_word_document(
    title: str,
    intro_text: Optional[str],
    images_data: List[Dict[str, Any]],
    page_orientation: str = "portrait",
    image_alignment: str = "center",
    image_size: str = "standard",
    one_image_per_page: bool = False
) -> Dict[str, Any]:
    """
    Generates a genuine, beautifully styled .docx document containing photos,
    captions, introductory text, and custom layout settings.
    """
    doc = Document()
    
    # 1. Configure Section & Page Setup
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    if page_orientation.lower() == "landscape":
        section.orientation = WD_ORIENTATION.LANDSCAPE
        section.page_width = Inches(11.0)
        section.page_height = Inches(8.5)
        printable_width = 11.0 - 1.6 # 9.4 inches
    else:
        section.orientation = WD_ORIENTATION.PORTRAIT
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        printable_width = 8.5 - 1.6 # 6.9 inches

    # Alignment enum mapping
    align_map = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT
    }
    alignment = align_map.get(image_alignment.lower(), WD_ALIGN_PARAGRAPH.CENTER)

    # Calculate image width in inches
    if image_size == "compact":
        img_target_width = min(4.2, printable_width)
    elif image_size == "large":
        img_target_width = min(6.4, printable_width)
    elif image_size == "full":
        img_target_width = printable_width
    else: # standard
        img_target_width = min(5.5, printable_width)

    # 2. Document Title
    clean_title = title.strip() if title and title.strip() else "Photo Collection Document"
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(4)
    title_p.paragraph_format.space_after = Pt(12)
    
    title_run = title_p.add_run(clean_title)
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(24)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(30, 58, 138) # Deep navy #1E3A8A

    # Subtle divider border/accent
    divider_p = doc.add_paragraph()
    divider_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    divider_p.paragraph_format.space_before = Pt(0)
    divider_p.paragraph_format.space_after = Pt(16)
    divider_run = divider_p.add_run("―" * 28)
    divider_run.font.size = Pt(10)
    divider_run.font.color.rgb = RGBColor(203, 213, 225) # Slate-300

    # 3. Optional Introductory Text
    if intro_text and intro_text.strip():
        intro_p = doc.add_paragraph()
        intro_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        intro_p.paragraph_format.space_after = Pt(20)
        intro_p.paragraph_format.line_spacing = 1.15
        
        intro_run = intro_p.add_run(intro_text.strip())
        intro_run.font.name = "Calibri"
        intro_run.font.size = Pt(11)
        intro_run.font.color.rgb = RGBColor(51, 65, 85) # Slate-700

    # 4. Insert Images with Captions
    processed_images_meta = []
    total_images = len(images_data)

    for index, item in enumerate(images_data):
        file_path = Path(item["filepath"])
        caption = item.get("caption", "").strip()

        if not file_path.exists():
            continue

        try:
            # Process & compress image with EXIF orientation
            image_buffer = fix_image_orientation_and_compress(file_path)

            # Insert Image paragraph
            img_p = doc.add_paragraph()
            img_p.alignment = alignment
            img_p.paragraph_format.space_before = Pt(14)
            img_p.paragraph_format.space_after = Pt(4)

            # Add picture
            img_p.add_run().add_picture(image_buffer, width=Inches(img_target_width))

            # Add Caption if present or image index
            cap_p = doc.add_paragraph()
            cap_p.alignment = alignment
            cap_p.paragraph_format.space_before = Pt(2)
            cap_p.paragraph_format.space_after = Pt(16)

            if caption:
                cap_run = cap_p.add_run(caption)
                cap_run.font.name = "Calibri"
                cap_run.font.size = Pt(10)
                cap_run.font.italic = True
                cap_run.font.color.rgb = RGBColor(100, 116, 139) # Slate-500
            else:
                # Subtle photo counter
                cap_run = cap_p.add_run(f"Figure {index + 1}")
                cap_run.font.name = "Calibri"
                cap_run.font.size = Pt(9)
                cap_run.font.italic = True
                cap_run.font.color.rgb = RGBColor(148, 163, 184)

            processed_images_meta.append({
                "original_filename": item.get("original_filename", file_path.name),
                "caption": caption,
                "order": index + 1
            })

            # Optional page break if requested and not the last image
            if one_image_per_page and (index < total_images - 1):
                doc.add_page_break()

        except Exception as e:
            print(f"Error processing image {file_path}: {e}")

    # 5. Header and Footer
    add_footer_and_header(doc, clean_title)

    # 6. Save document
    doc_id = str(uuid.uuid4())
    sanitized = sanitize_filename(clean_title)
    final_filename = f"{sanitized}.docx"
    
    # Ensure unique filename on disk if collision
    disk_filename = f"{sanitized}_{doc_id[:8]}.docx"
    save_path = DOCS_DIR / disk_filename

    doc.save(str(save_path))
    file_size_bytes = save_path.stat().st_size

    return {
        "id": doc_id,
        "title": clean_title,
        "filename": final_filename,
        "disk_filename": disk_filename,
        "filepath": str(save_path),
        "image_count": len(processed_images_meta),
        "file_size_bytes": file_size_bytes,
        "page_orientation": page_orientation,
        "image_alignment": image_alignment,
        "image_size": image_size,
        "intro_text": intro_text or "",
        "images_metadata": processed_images_meta
    }
